from __future__ import annotations

import csv
import json
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Protocol

from openpyxl import load_workbook

from .config import Settings
from .mws_client import MwsClient
from .utils import normalize_name


@dataclass(frozen=True)
class RomanianPlayer:
    name: str
    normalized_name: str
    team: str
    position: str
    market_value_eur: int | None
    source_url: str


@dataclass(frozen=True)
class AuctionMatch:
    checked_at: str
    player_name: str
    transfermarkt_team: str
    position: str
    market_value_eur: int | None
    mws_name: str
    event_name: str
    labels: str
    current_bid_eur: int | None
    current_bid_usd: int | None
    end_date_utc: str
    slug: str
    url: str


class ProductClient(Protocol):
    def get_json(self, path: str, params: dict[str, Any] | None = None, cache: bool = True) -> dict[str, Any]:
        ...


def load_romanian_players(path: str | Path) -> list[RomanianPlayer]:
    source = Path(path)
    if source.suffix.lower() == ".csv":
        return _load_players_from_csv(source)
    if source.suffix.lower() in {".xlsx", ".xlsm"}:
        return _load_players_from_xlsx(source)
    raise ValueError(f"Unsupported player list format: {source.suffix}")


def _load_players_from_csv(path: Path) -> list[RomanianPlayer]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        return [_player_from_row(row) for row in reader if row.get("Nume") or row.get("name")]


def _load_players_from_xlsx(path: Path) -> list[RomanianPlayer]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Jucători"] if "Jucători" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    header_index = next((i for i, row in enumerate(rows) if "Nume" in [str(cell) for cell in row if cell is not None]), None)
    if header_index is None:
        raise ValueError("Could not find a header row containing 'Nume'")
    headers = [str(cell) if cell is not None else "" for cell in rows[header_index]]
    players = []
    for raw in rows[header_index + 1 :]:
        row = {headers[i]: raw[i] for i in range(min(len(headers), len(raw))) if headers[i]}
        if row.get("Nume"):
            players.append(_player_from_row(row))
    return players


def _player_from_row(row: dict[str, Any]) -> RomanianPlayer:
    name = str(row.get("Nume") or row.get("name") or "").strip()
    value = row.get("M.V. (EUR)") or row.get("market_value_eur") or row.get("value")
    market_value = _int_or_none(value)
    return RomanianPlayer(
        name=name,
        normalized_name=normalize_name(name),
        team=str(row.get("Echipă") or row.get("team") or "").strip(),
        position=str(row.get("Post") or row.get("position") or "").strip(),
        market_value_eur=market_value,
        source_url=str(row.get("Profil Transfermarkt") or row.get("source_url") or "").strip(),
    )


def _int_or_none(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(float(str(value).replace(",", "")))
    except ValueError:
        return None


def fetch_live_bidding_products(client: ProductClient, page_size: int = 120) -> list[dict[str, Any]]:
    products: list[dict[str, Any]] = []
    offset = 0
    while True:
        data = client.get_json(
            "/v2/products/paginated",
            {
                "only_return_not_ended": True,
                "sort_products_by": "closing_time_first_to_last",
                "offset": offset,
                "page_size": page_size,
            },
            cache=False,
        )
        page = data.get("results") or []
        products.extend(page)
        if not page:
            break
        total = int(data.get("total_count") or 0)
        offset += page_size
        if total and offset >= total:
            break
    return products


def match_players_to_products(
    players: list[RomanianPlayer],
    products: list[dict[str, Any]],
    checked_at_iso: str | None = None,
    web_base_url: str = "https://mws.com/us/product",
) -> list[AuctionMatch]:
    checked_at = checked_at_iso or datetime.now(timezone.utc).isoformat()
    players_by_name = {player.normalized_name: player for player in players}
    matches = []
    seen_slugs: set[str] = set()
    for product in products:
        if product.get("sales_method") != "bidding":
            continue
        if product.get("finished") is True:
            continue
        if product.get("type") != "football_shirt":
            continue
        player = players_by_name.get(normalize_name(product.get("name") or product.get("canonical_name")))
        if not player:
            continue
        slug = product.get("slug") or ""
        if slug in seen_slugs:
            continue
        seen_slugs.add(slug)
        prices = {row.get("currency"): row.get("price") for row in product.get("prices") or [] if isinstance(row, dict)}
        matches.append(
            AuctionMatch(
                checked_at=checked_at,
                player_name=player.name,
                transfermarkt_team=player.team,
                position=player.position,
                market_value_eur=player.market_value_eur,
                mws_name=product.get("name") or "",
                event_name=product.get("product_event_name") or product.get("event_name") or "",
                labels=_labels_text(product),
                current_bid_eur=prices.get("EUR"),
                current_bid_usd=prices.get("USD"),
                end_date_utc=product.get("end_date") or "",
                slug=slug,
                url=f"{web_base_url.rstrip('/')}/{slug}",
            )
        )
    return sorted(matches, key=lambda item: (item.end_date_utc, item.player_name, item.slug))


def _labels_text(product: dict[str, Any]) -> str:
    labels = []
    for label in product.get("labels") or []:
        if isinstance(label, dict):
            labels.append(label.get("text") or "")
        else:
            labels.append(str(label))
    return ", ".join(label for label in labels if label)


def write_outputs(matches: list[AuctionMatch], output_dir: str | Path, checked_at_iso: str) -> None:
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    rows = [asdict(match) for match in matches]
    fields = list(asdict(matches[0]).keys()) if matches else list(AuctionMatch.__dataclass_fields__.keys())
    with (output / "latest.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
    with (output / "latest.json").open("w", encoding="utf-8") as handle:
        json.dump({"checked_at": checked_at_iso, "match_count": len(matches), "matches": rows}, handle, ensure_ascii=False, indent=2)
    (output / "latest.md").write_text(_markdown(matches, checked_at_iso), encoding="utf-8")


def _markdown(matches: list[AuctionMatch], checked_at_iso: str) -> str:
    lines = [
        "# Romanian MWS Auction Matches",
        "",
        f"Checked at: {checked_at_iso}",
        "",
        "| Player | Item | Event | Bid EUR | Bid USD | Ends UTC |",
        "|---|---|---|---:|---:|---|",
    ]
    for match in matches:
        lines.append(
            f"| {match.player_name} | [{match.labels or 'Auction'}]({match.url}) | {match.event_name} | "
            f"{match.current_bid_eur or ''} | {match.current_bid_usd or ''} | {match.end_date_utc} |"
        )
    if not matches:
        lines.append("| No matches |  |  |  |  |  |")
    lines.append("")
    return "\n".join(lines)


def run_romanian_auctions(
    player_list_path: str | Path,
    output_dir: str | Path,
    settings: Settings,
) -> list[AuctionMatch]:
    checked_at = datetime.now(timezone.utc).isoformat()
    players = load_romanian_players(player_list_path)
    client = MwsClient(settings)
    try:
        products = fetch_live_bidding_products(client)
    finally:
        client.close()
    matches = match_players_to_products(players, products, checked_at_iso=checked_at, web_base_url=settings.mws_web_base_url)
    write_outputs(matches, output_dir, checked_at)
    return matches
