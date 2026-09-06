from __future__ import annotations

import csv
import math
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .utils import normalize_name


@dataclass(frozen=True)
class CompRow:
    market_value_eur: float
    last_eur: float
    type_bucket: str
    fair_eur: float


def labels_to_mws_type(labels: list[dict[str, Any]] | None) -> str:
    texts: list[str] = []
    for lab in labels or []:
        if isinstance(lab, dict) and lab.get("text"):
            texts.append(str(lab["text"]))
    return ", ".join(texts)


def mws_type_bucket(mws_type: str) -> str:
    t = (mws_type or "").lower()
    if "worn" in t and "signed" in t:
        return "worn_signed"
    if "worn" in t:
        return "worn"
    if "issued" in t and "signed" in t:
        return "issued_signed"
    if "issued" in t:
        return "issued"
    return "other"


def bid_ladder(fair_eur: int) -> tuple[int, int, int]:
    """Match the workbook convention: steal 60%, strong 75%, max 85% of fair (rounded)."""
    if fair_eur <= 0:
        return 0, 0, 0
    steal = int(round(0.6 * fair_eur))
    strong = int(round(0.75 * fair_eur))
    max_bid = int(round(0.85 * fair_eur))
    return steal, strong, max_bid


def _log1p_eur(x: float) -> float:
    return math.log10(max(x, 0.0) + 1.0)


def load_comp_rows_from_workbook(path: str | Path, sheet: str = "Ranked All 200") -> list[CompRow]:
    path = Path(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
    except KeyError as exc:
        wb.close()
        raise ValueError(f"Sheet {sheet!r} not found in {path}") from exc
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows: list[CompRow] = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(headers, r))
        mv = d.get("market_value_eur")
        last = d.get("mws_last_price_eur")
        fair = d.get("fair_price_eur")
        mtype = str(d.get("mws_type") or "")
        if mv is None or last is None or fair is None:
            continue
        try:
            mv_f = float(mv)
            last_f = float(last)
            fair_f = float(fair)
        except (TypeError, ValueError):
            continue
        if mv_f <= 0 or last_f <= 0 or fair_f <= 0:
            continue
        rows.append(CompRow(mv_f, last_f, mws_type_bucket(mtype), fair_f))
    wb.close()
    return rows


def estimate_fair_knn(
    market_value_eur: float,
    last_price_eur: float,
    mws_type: str,
    comps: list[CompRow],
    *,
    k: int = 7,
) -> tuple[int, str]:
    """Fair price EUR from k nearest comps in (log market, log last) space, preferring the same type bucket."""
    if market_value_eur <= 0 or last_price_eur <= 0:
        return 0, "missing market_value_eur or last comp price"
    if not comps:
        return 0, "no comparable rows in workbook"
    tb = mws_type_bucket(mws_type)
    pool_same = [c for c in comps if c.type_bucket == tb]
    pool = pool_same if len(pool_same) >= max(4, k // 2) else list(comps)
    bucket_note = "same-type" if pool is pool_same else "all-types"
    qx = _log1p_eur(market_value_eur)
    qy = _log1p_eur(last_price_eur)
    scored: list[tuple[float, float]] = []
    for c in pool:
        dx = _log1p_eur(c.market_value_eur) - qx
        dy = _log1p_eur(c.last_eur) - qy
        scored.append((math.hypot(dx, dy), c.fair_eur))
    scored.sort(key=lambda t: t[0])
    fairs = [fair for _, fair in scored[:k]]
    fairs.sort()
    if not fairs:
        return 0, "no neighbors"
    mid = len(fairs) // 2
    if len(fairs) % 2:
        med = fairs[mid]
    else:
        med = 0.5 * (fairs[mid - 1] + fairs[mid])
    return int(round(med)), f"kNN k={len(fairs)} {bucket_note} bucket={tb}"


def load_tm_overrides_csv(path: str | Path) -> tuple[dict[str, dict[str, int | None]], dict[str, dict[str, int | None]]]:
    """
    CSV columns: name, market_value_eur, optional mws_last_price_eur, optional mws_slug.

    Rows are keyed by normalized player name. Optional mws_slug (full product slug from MWS)
    registers the same row under that lowercase slug for disambiguation (e.g. short \"Kevin\").
    """
    by_name: dict[str, dict[str, int | None]] = {}
    by_slug: dict[str, dict[str, int | None]] = {}
    with Path(path).open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames or "name" not in reader.fieldnames:
            raise ValueError("tm-csv must include a 'name' column (may be blank if mws_slug is set)")
        for row in reader:
            raw_name = (row.get("name") or "").strip()
            slug = (row.get("mws_slug") or "").strip().lower()
            if not raw_name and not slug:
                continue
            mv = (row.get("market_value_eur") or "").strip().replace(",", "")
            last = (row.get("mws_last_price_eur") or "").strip().replace(",", "")
            payload: dict[str, int | None] = {
                "market_value_eur": int(mv) if mv else None,
                "mws_last_price_eur": int(last) if last else None,
            }
            if raw_name:
                by_name[normalize_name(raw_name)] = payload
            if slug:
                by_slug[slug] = payload
    return by_name, by_slug


def parse_event_player_title(canonical_name: str) -> tuple[str, str | None]:
    """'Name | Club' -> (player_name, club)"""
    text = (canonical_name or "").strip()
    if "|" in text:
        left, right = text.split("|", 1)
        return left.strip(), right.strip() or None
    return text, None
